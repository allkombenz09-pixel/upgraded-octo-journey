#!/usr/bin/env python3
"""Build the production package (.docx, Russian) for the voice/image helper
from episode.json, so the document and the edit can never drift apart.

  python tools/build_package.py episodes/01-ikea
"""
import copy
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

TEMPLATE = Path(__file__).resolve().parent / "filmkit" / "assets" / "package_template.docx"
WPS = 2.3  # spoken words per second at the channel's voice settings, pauses included
ACCENT = RGBColor(0xC4, 0x62, 0x3F)
MUTED = RGBColor(0x6B, 0x72, 0x80)


class Doc:
    def __init__(self):
        self.d = Document(str(TEMPLATE))
        self.numbering = self.d.part.numbering_part.element
        self.next_num = 100

    def _num(self, abstract):
        """New list instance (restarts at 1). abstract 0 = bullets, 1 = decimal."""
        self.next_num += 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(self.next_num))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), str(abstract))
        num.append(a)
        if abstract == 1:
            o = OxmlElement("w:lvlOverride")
            o.set(qn("w:ilvl"), "0")
            s = OxmlElement("w:startOverride")
            s.set(qn("w:val"), "1")
            o.append(s)
            num.append(o)
        self.numbering.append(num)
        return self.next_num

    def runs(self, p, text, size=None, color=None):
        """**bold**, _italic_ and {{tag}} (accent) markup."""
        for part in re.split(r"(\*\*.+?\*\*|\{\{.+?\}\}|_.+?_)", text):
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("{{"):
                r = p.add_run(part[2:-2])
                r.bold = True
                r.font.color.rgb = ACCENT
            elif part.startswith("_") and part.endswith("_") and len(part) > 2:
                r = p.add_run(part[1:-1])
                r.italic = True
            else:
                r = p.add_run(part)
            if size:
                r.font.size = Pt(size)
            if color is not None and not part.startswith("{{"):
                r.font.color.rgb = color
        return p

    def h1(self, t):
        return self.d.add_paragraph(t, style="Heading 1")

    def h2(self, t):
        return self.d.add_paragraph(t, style="Heading 2")

    def h3(self, t):
        return self.d.add_paragraph(t, style="Heading 3")

    def p(self, t="", size=None, color=None):
        return self.runs(self.d.add_paragraph(style="Normal"), t, size, color)

    def code(self, text):
        p = self.d.add_paragraph(style="Code")
        lines = text.split("\n")
        for i, line in enumerate(lines):
            r = p.add_run(line)
            if i < len(lines) - 1:
                r.add_break()
        self.d.add_paragraph(style="Normal")
        return p

    def lst(self, items, numbered=False):
        nid = self._num(1 if numbered else 0)
        for it in items:
            p = self.d.add_paragraph(style="List Paragraph")
            pPr = p._p.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            il = OxmlElement("w:ilvl")
            il.set(qn("w:val"), "0")
            ni = OxmlElement("w:numId")
            ni.set(qn("w:val"), str(nid))
            numPr.append(il)
            numPr.append(ni)
            pPr.append(numPr)
            self.runs(p, it)

    def checklist(self, items):
        for it in items:
            p = self.d.add_paragraph(style="Normal")
            p.paragraph_format.space_after = Pt(2)
            self.runs(p, "☐  " + it)
        self.d.add_paragraph(style="Normal")

    def table(self, rows, widths=None, size=None):
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ""
                self.runs(cell.paragraphs[0], ("**" + str(val) + "**") if i == 0 and val else str(val), size)
        self.d.add_paragraph(style="Normal")
        return t

    def save(self, path):
        self.d.save(str(path))


def words(text):
    return len(text.split())


def mmss(sec):
    return f"{int(sec // 60)}:{int(round(sec % 60)):02d}"


def strip_tail(prompt, tail):
    return prompt[: -len(tail)].rstrip(", ") if tail and prompt.endswith(tail) else prompt


def build(ep_dir):
    ep = json.loads((ep_dir / "episode.json").read_text())
    pk = ep.get("package_ru", {})
    tail = ep.get("style_tail", "")
    images = [s for s in ep["shots"] if s.get("kind") == "image"]
    new = [s["id"] for s in images if s.get("status") == "new"]
    edited = [s["id"] for s in images if s.get("status") == "edited"]
    kling = [s["id"] for s in ep["shots"] if s.get("kling")]
    main_words = sum(words(b["text"]) for p in ep["parts"] for b in p["beats"])

    D = Doc()
    D.h1(pk.get("title", ep.get("title", "Пакет производства")))
    D.p(pk.get("byline", ""), color=MUTED)

    if pk.get("changes"):
        D.h2("Что изменилось в v2")
        D.lst(pk["changes"])

    D.h2("Порядок работы")
    D.p(pk.get("intro", ""))
    D.lst(pk.get("steps", []), numbered=True)
    D.p(pk.get("after", ""))
    D.h3("Чек-лист")
    D.checklist(pk.get("checklist", []))
    if pk.get("naming"):
        D.p("**Как называть файлы** (папка с выпуском, без подпапок):")
        D.table([["Что", "Имя файла"]] + pk["naming"])

    D.h2("Настройки стиля и голоса")
    D.p("Рисованный стиль и один голос — одинаково во всех роликах канала.")
    D.p("Хвост стиля — **уже вставлен** в конец каждого промпта ниже, отдельно добавлять не нужно:")
    D.code(tail)
    D.table([["Параметр", "Значение"]] + pk.get("settings", []))
    for line in pk.get("style_notes", []):
        D.p(line)

    D.h2("Текст озвучки")
    D.p(pk.get("voice_intro", "").format(minutes=round(main_words / WPS / 60, 1), words=main_words))
    for part in ep["parts"]:
        text = "\n\n".join(b["text"] for b in part["beats"])
        D.p(f"**{part['title_ru']}** → файл {{{{{part['id']}.mp3}}}} · {words(text)} слов · ≈ {mmss(words(text) / WPS)}")
        D.code(text)

    D.h2(f"Промпты картинок ({len(images)} шт.)")
    D.p(pk.get("prompts_intro", ""))
    current = None
    beat_part = {b["id"]: p["title_ru"] for p in ep["parts"] for b in p["beats"]}
    for s in images:
        part = beat_part.get((s.get("beats") or [""])[0], current)
        if part != current:
            D.h3(part.split(" — ")[0] if part else "")
            current = part
        tag = {"new": " {{НОВЫЙ}}", "edited": " {{ИЗМЕНЁН}}"}.get(s.get("status"), "")
        extra = " · {{Kling}}" if s.get("kling") else ""
        para = D.p(f"**{s['id']}**{tag}{extra} — {s['prompt']}")
        para.paragraph_format.space_after = Pt(4)
        if s.get("ref"):
            D.p(f"_{s['ref']}_", color=MUTED)
    D.h3("Обложка")
    D.p(f"**thumbnail** — {ep['packaging']['thumbnail_prompt']}")
    D.p(pk.get("thumbnail_note", ""), color=MUTED)

    if kling:
        D.h2("Kling (по желанию)")
        D.p(pk.get("kling_intro", "").format(ids=", ".join(kling)))
        D.code(ep.get("kling_prompt", ""))

    D.h2("4 Shorts")
    D.p(pk.get("shorts_intro", ""))
    rows = [["Short", "Тема", "Кадры", "Слов / ≈ сек"]]
    for sh in ep["shorts"]:
        w = sum(words(b["text"]) for b in sh["beats"])
        rows.append([sh["id"].replace("short", ""), sh["title_ru"], ", ".join(c["shot"] for c in sh["cuts"]), f"{w} / ≈ {round(w / WPS)}"])
    D.table(rows)
    for sh in ep["shorts"]:
        D.p(f"**Short {sh['id'][-1]} — {sh['title_ru']}** → файл {{{{{sh['id']}.mp3}}}}")
        D.code("\n\n".join(b["text"] for b in sh["beats"]))

    D.h2("Упаковка")
    pkg = ep["packaging"]
    D.p("Основной заголовок — вариант 1; 2 и 3 — для теста в YouTube «Test & Compare».")
    D.lst(pkg["titles"], numbered=True)
    D.p(f"Текст на обложке (добавлю я): **{pkg['thumbnail_text']}**")
    D.p("**Описание** (главы подставляются автоматически из монтажа — файл out/description.txt):")
    D.code(pkg["description_body"])
    D.p("**При загрузке:**")
    D.checklist(pk.get("upload_checklist", []))

    if pk.get("facts_note") or ep.get("sources"):
        D.h2("Проверка фактов")
        D.p(pk.get("facts_note", ""))
        if ep.get("sources"):
            D.lst(ep["sources"])

    out = ep_dir / "package" / pk.get("filename", "production_v2.docx")
    D.d.core_properties.title = pk.get("title", "")
    D.save(out)
    return out


if __name__ == "__main__":
    print(build(Path(sys.argv[1])))
